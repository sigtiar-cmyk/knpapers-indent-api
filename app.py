from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import json, io, base64, os

app = Flask(__name__)

# Allow ALL origins - needed for Netlify to call this API
CORS(app, origins="*", methods=["GET","POST","OPTIONS"], 
     allow_headers=["Content-Type","Authorization"])

TEMPLATE_B64 = open('/opt/render/project/src/template_b64.txt').read().strip()

@app.route('/', methods=['GET'])
def home():
    return jsonify({'status': 'KN Papers Indent API is running!', 'version': '1.0'})

@app.route('/health', methods=['GET'])
def health():
    return jsonify({'status': 'ok', 'service': 'KN Papers Indent API'})

@app.route('/generate-indent', methods=['POST', 'OPTIONS'])
def generate_indent():
    # Handle preflight CORS request
    if request.method == 'OPTIONS':
        response = jsonify({'status': 'ok'})
        response.headers['Access-Control-Allow-Origin'] = '*'
        response.headers['Access-Control-Allow-Methods'] = 'POST, OPTIONS'
        response.headers['Access-Control-Allow-Headers'] = 'Content-Type'
        return response

    try:
        from openpyxl import load_workbook
        
        data = request.json
        template_bytes = base64.b64decode(TEMPLATE_B64)
        template_io = io.BytesIO(template_bytes)
        wb = load_workbook(template_io)
        ws = wb.active

        def W(cell, value):
            try: ws[cell] = value
            except: pass

        iNo    = str(data.get('iNo', ''))
        iMonth = str(data.get('iMonth', 'MAY')).upper()
        iDate  = str(data.get('iDate', ''))
        d      = data.get('dealer') or {}
        c      = data.get('consignee') or d
        b      = data.get('buyer') or d
        items  = data.get('items', [])
        trans  = data.get('trans', 'ROAD')
        ins    = data.get('ins', 'BY CONSIGNEE')
        deliv  = data.get('deliv', 'IMMEDIATE/ STANDARD')

        # Indent header
        W('C4', iMonth); W('D4', iNo); W('C5', iDate)

        # Dealer
        W('C7', d.get('name','')); W('C8', d.get('addr',''))
        W('C9', 'Ph No: ' + str(d.get('ph','')))
        W('B11', d.get('gstin','')); W('F11', d.get('st',''))

        # Consignee
        W('D13', c.get('name', d.get('name','')))
        W('C14', c.get('addr', d.get('addr','')))
        W('C15', 'Ph No: ' + str(c.get('ph', d.get('ph',''))))
        W('B18', c.get('gstin', d.get('gstin',''))); W('F18', c.get('st', d.get('st','')))

        # Buyer
        W('C20', b.get('name', d.get('name','')))
        W('C21', b.get('addr', d.get('addr','')))
        W('C22', 'Ph No: ' + str(b.get('ph', d.get('ph',''))))
        W('B24', b.get('gstin', d.get('gstin',''))); W('F24', b.get('st', d.get('st','')))

        # Transport & Insurance
        W('L23', trans)
        if ins == 'SELF':
            W('D26', 'SELF  \u2714'); W('J26', '')
        else:
            W('D26', 'SELF'); W('J26', '\u2714')

        # Items
        for i in range(8):
            r = 30 + i
            it = items[i] if i < len(items) else {}
            W(f'A{r}', it.get('sl', i+1))
            W(f'B{r}', it.get('mc', ''))
            W(f'C{r}', it.get('code', ''))
            W(f'D{r}', it.get('prod', ''))
            W(f'F{r}', it.get('hsn', ''))
            W(f'G{r}', it.get('gsm', ''))
            W(f'H{r}', it.get('sz', ''))
            try: W(f'I{r}', float(it['rw']) if it.get('rw') else '')
            except: W(f'I{r}', '')
            W(f'J{r}', it.get('ns', ''))
            W(f'K{r}', it.get('col', ''))
            try: W(f'L{r}', float(it['qty']) if it.get('qty') else '')
            except: W(f'L{r}', '')
            W(f'M{r}', it.get('deck', ''))
            W(f'N{r}', it.get('rem', ''))

        # Instructions & discounts
        W('D41', deliv)
        W('L41', 'ALLOW / DISALLOW')
        W('L42', 'ALLOW / DISALLOW/NOT APPLICABLE')
        W('L43', 'ALLOW / DISALLOW/NOT APPLICABLE')
        W('L44', 'ALLOW / DISALLOW/NOT APPLICABLE')
        W('L45', 'ALLOW / DISALLOW')
        W('K49', 'For: ' + d.get('name', 'SHRI K.N.PAPERS'))

        # Save
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'WCPM_Indent_{iMonth}_{iNo}.xlsx'
        )
        response.headers['Access-Control-Allow-Origin'] = '*'
        return response

    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
